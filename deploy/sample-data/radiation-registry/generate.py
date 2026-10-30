"""Builds the radiation-registry sample packages from the legacy system's Excel exports.

Usage (from the repository root):
    python3 deploy/sample-data/radiation-registry/generate.py LOOKUPS.xlsx ROLES.xlsx FACILITIES.xlsx

Needs openpyxl. Writes the numbered JSON packages next to this script; import them in that order.
"""

import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent
MAX_ITEMS = 500  # the import endpoint takes at most 500 items per package
LEGACY = "LEGACY"  # external-code provider holding the legacy system's term ids
GENERAL_CODE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.\-]{0,99}$")

# Classes whose sets carry no Arabic name in the export.
CLASS_NAMES_AR = {"Province": "المنطقة", "Region": "المنطقة الإدارية"}
# Typos in the export's Arabic names, fixed on the way in.
ARABIC_FIXES = {"لمنطقة الشرقية": "المنطقة الشرقية"}
# Extra data columns that mean something for a class become typed attributes; the rest are dropped.
CLASS_ATTRIBUTES = {
    "ActivityUnit": [("BQ_FACTOR", "معامل التحويل إلى بيكريل", "Factor to becquerel", "NUMBER", 1)],
    "Status": [
        ("EXTRA_DATA_1", "بيانات إضافية 1", "Extra data 1", "TEXT", 1),
        ("EXTRA_DATA_2", "بيانات إضافية 2", "Extra data 2", "TEXT", 2),
    ],
}
# Permission codes of the legacy system → privilege actions (standard ones where they mean the same).
STANDARD_ACTIONS = {"view": "VIEW", "add": "CREATE", "edit": "EDIT", "delete": "DELETE", "export": "EXPORT", "print": "PRINT"}
ACTION_FIXES = {
    "viewsourcesbutton": "VIEW_SOURCES_BUTTON",
    "viewrepresntativesbutton": "VIEW_REPRESENTATIVES_BUTTON",
    "viewlicensebutton": "VIEW_LICENSE_BUTTON",
    "viewnmbutton": "VIEW_NM_BUTTON",
    "viewpermitbutton": "VIEW_PERMIT_BUTTON",
    "viewworkerbutton": "VIEW_WORKER_BUTTON",
}
ADMIN_ROLE_AR = "مسؤول النظام"
# Cities of the facilities export, under the Region lookup's values.
CITIES = [
    ("AL_AHSA", "الأحساء", "Al Ahsa", "5"),
    ("YANBU", "ينبع", "Yanbu", "3"),
    ("AL_JUBAYL", "الجبيل", "Al Jubayl", "5"),
    ("AR_RIYAD", "الرياض", "Ar Riyad", "1"),
    ("AD_DAMMAM", "الدمام", "Ad Dammam", "5"),
    ("JIDDAH", "جدة", "Jiddah", "2"),
]


def upper_snake(name):
    return re.sub(r"(?<=[a-z0-9])(?=[A-Z])", "_", name).replace("-", "_").upper()


def text(value):
    if value is None:
        return None
    value = str(value).strip()
    value = re.sub(r"\s+", " ", value)
    return value or None


def arabic(value):
    value = text(value)
    return ARABIC_FIXES.get(value, value)


def item(entity_type, code, payload):
    return {"entityType": entity_type, "entityCode": code, "operation": "CREATE", "payload": payload}


def rows(ws):
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in row):
            yield dict(zip(header, row))


def lookups(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    sets = list(rows(wb["Lookup sets"]))
    terms = list(rows(wb["Lookup terms"]))
    classes, values, notes = [], OrderedDict(), []
    for s in sets:
        name = s["Class name"]
        code = upper_snake(name)
        name_ar = CLASS_NAMES_AR.get(name, text(s["Name (Arabic)"]))
        attributes = [
            {"code": c, "nameAr": ar, "nameEn": en, "dataType": t, "displayOrder": order * 10}
            for c, ar, en, t, order in CLASS_ATTRIBUTES.get(name, [])
        ]
        classes.append(item("apps.class", code, {
            "nameAr": name_ar,
            "nameEn": text(s["Name (English)"]),
            "isActive": s["Active"] == "Yes",
            "descriptionEn": f"Imported from the legacy lookup set {name} (id {s['Id']}).",
            "attributes": attributes,
        }))
        class_terms = [t for t in terms if t["Class name"] == name]
        seen_codes, seen_names = set(), set()
        # Short legacy values are kept as codes when they are valid and unique; otherwise the term id is the code.
        candidate = [text(t["Value"]) for t in class_terms]
        usable = all(v and GENERAL_CODE.match(v) and len(v) <= 20 for v in candidate) and len(set(candidate)) == len(candidate)
        out = []
        for order, t in enumerate(class_terms, start=1):
            name_en = text(t["Name (English)"]) or text(t["Value"])
            if not name_en:
                notes.append(f"{name}: skipped term {t['Term id']} (no value or name)")
                continue
            key = name_en.casefold()
            if key in seen_names:
                notes.append(f"{name}: skipped duplicate '{name_en}' (term {t['Term id']})")
                continue
            seen_names.add(key)
            value_code = text(t["Value"]) if usable else str(t["Term id"])
            seen_codes.add(value_code)
            attrs = {}
            for attr_code, _, _, data_type, column in CLASS_ATTRIBUTES.get(name, []):
                raw = text(t[f"Extra data {column}"])
                if raw is None or raw == "NULL":
                    continue
                attrs[attr_code] = float(raw) if data_type == "NUMBER" else raw
                if data_type == "NUMBER" and attrs[attr_code].is_integer():
                    attrs[attr_code] = int(attrs[attr_code])
            payload = {
                "nameAr": arabic(t["Name (Arabic)"]) or name_en,
                "nameEn": name_en,
                "isActive": t["Active"] == "Yes",
                "displayOrder": order * 10,
                "externalCodes": [{"providerCode": LEGACY, "externalCode": str(t["Term id"]), "direction": "BOTH"}],
            }
            if attrs:
                payload["attributes"] = attrs
            out.append(item("apps.domain", f"{code}/{value_code}", payload))
        values[code] = out
    return classes, values, notes


def action_code(permission):
    if permission in STANDARD_ACTIONS:
        return STANDARD_ACTIONS[permission]
    if permission in ACTION_FIXES:
        return ACTION_FIXES[permission]
    return upper_snake(permission)


def security(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    screens = OrderedDict()
    for r in rows(wb["Privileges and permissions"]):
        screen = screens.setdefault(r["Screen (privilege)"], {"en": text(r["Screen (English)"]), "ar": text(r["Screen (Arabic)"]), "actions": []})
        code = action_code(r["Permission"])
        if code not in screen["actions"]:
            screen["actions"].append(code)
    resources = [
        item("apps.resource", code, {
            "nameAr": s["ar"], "nameEn": s["en"], "appCode": "INTERNAL", "resourceType": "SCREEN", "availableActions": s["actions"],
        })
        for code, s in screens.items()
    ]

    roles_by_name = {}
    roles = []
    grants = OrderedDict()
    for r in rows(wb["Role permissions"]):
        actions = grants.setdefault(r["Role"], OrderedDict()).setdefault(r["Screen (privilege)"], [])
        code = action_code(r["Permission"])
        if code not in actions:
            actions.append(code)
    for r in rows(wb["Roles"]):
        name_en = text(r["Name (English)"])
        name_ar = text(r["Name (Arabic)"])
        if name_ar == name_en:
            name_ar = ADMIN_ROLE_AR
        roles_by_name[name_en] = r["Code"]
        privileges = [{"resourceCode": screen, "actions": actions} for screen, actions in grants.get(name_en, {}).items()]
        roles.append(item("apps.role", r["Code"], {
            "nameAr": name_ar, "nameEn": name_en, "appCode": "INTERNAL", "referenceCode": str(r["Id"]), "privileges": privileges,
        }))

    users = OrderedDict()
    for r in rows(wb["Role users"]):
        user = users.setdefault(r["User name"], {"name": text(r["Display name"]), "email": text(r["Email"]), "active": r["Active"] == "Yes", "roles": []})
        role = roles_by_name[r["Role"]]
        if role not in user["roles"]:
            user["roles"].append(role)
    user_items = [
        item("apps.user", name, {
            "nameAr": u["name"], "nameEn": u["name"], "email": u["email"], "status": "ACTIVE" if u["active"] else "INACTIVE",
            "roles": [{"roleCode": r, "companyCode": "*", "branches": ["*"]} for r in u["roles"]],
        })
        for name, u in users.items()
    ]
    return resources + roles + user_items


def facilities(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    city_codes = {en: code for code, _, en, _ in CITIES}
    items = [item("apps.class", "CITY", {
        "nameAr": "المدينة", "nameEn": "City", "parentClassCode": "REGION",
        "descriptionEn": "Cities of the sample facilities, under the Region lookup.",
    })]
    items += [
        item("apps.domain", f"CITY/{code}", {"nameAr": ar, "nameEn": en, "parentCode": region, "displayOrder": i * 10})
        for i, (code, ar, en, region) in enumerate(CITIES, start=1)
    ]
    attributes = [
        ("CITY", "المدينة", "City", "DOMAIN_REF", True),
        ("NATIONAL_ADDRESS", "العنوان الوطني", "National address", "TEXT", False),
        ("PHONE", "هاتف المنشأة", "Facility phone", "TEXT", False),
        ("EMAIL", "بريد المنشأة", "Facility email", "TEXT", False),
        ("OPERATIONAL", "قيد التشغيل", "Operational", "BOOL", True),
        ("AMAN_ID", "رقم المنشأة في نظام أمان", "Facility AMAN ID", "TEXT", False),
        ("LICENCES", "عدد التراخيص", "Licences", "NUMBER", False),
    ]
    items.append(item("apps.class", "FACILITY", {
        "nameAr": "المنشأة", "nameEn": "Facility",
        "descriptionEn": "Licensee facilities exported from AMAN.",
        "attributes": [
            {"code": c, "nameAr": ar, "nameEn": en, "dataType": t, "isRequired": req, "displayOrder": i * 10,
             **({"refClassCode": "CITY"} if t == "DOMAIN_REF" else {})}
            for i, (c, ar, en, t, req) in enumerate(attributes, start=1)
        ],
    }))
    for i, r in enumerate(rows(wb["Facilities"]), start=1):
        attrs = {
            "CITY": city_codes[text(r["City"])],
            "NATIONAL_ADDRESS": text(r["National address"]),
            "PHONE": text(r["Facility phone"]),
            "EMAIL": text(r["Facility email"]),
            "OPERATIONAL": text(r["Operational status"]) == "Operational",
            "AMAN_ID": text(r["Facility AMAN ID"]),
            "LICENCES": r["Licences"],
        }
        name = text(r["Facility"])
        items.append(item("apps.domain", f"FACILITY/{text(r['Facility code'])}", {
            "nameAr": name, "nameEn": name, "displayOrder": i * 10,
            "attributes": {k: v for k, v in attrs.items() if v is not None},
        }))
    return items


def write(name, title, items):
    if len(items) > MAX_ITEMS:
        raise SystemExit(f"{name}: {len(items)} items, more than {MAX_ITEMS}")
    path = OUT / name
    path.write_text(json.dumps({"title": title, "items": items}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{name}: {len(items)} items")


def main():
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    classes, values, notes = lookups(sys.argv[1])
    manufacturers = values.pop("MANUFACTURER")
    write("01-lookups.json", "Radiation registry: lookups", classes + [v for vs in values.values() for v in vs])
    parts = (len(manufacturers) + MAX_ITEMS - 1) // MAX_ITEMS
    size = (len(manufacturers) + parts - 1) // parts
    for p in range(parts):
        write(f"02-manufacturers-{p + 1}.json", f"Radiation registry: manufacturers ({p + 1} of {parts})", manufacturers[p * size:(p + 1) * size])
    write("03-roles-and-permissions.json", "Radiation registry: screens, roles and users", security(sys.argv[2]))
    write("04-facilities.json", "Radiation registry: cities and facilities", facilities(sys.argv[3]))
    for note in notes:
        print("note:", note)


if __name__ == "__main__":
    main()
